from __future__ import annotations

import io
import os
import re
from copy import copy
from pathlib import Path
from urllib.parse import urljoin
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import (
    AnchorMarker,
    OneCellAnchor,
    TwoCellAnchor,
    XDRPositiveSize2D,
)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU, points_to_pixels
from PIL import Image as PillowImage
from core.logger import logger


# 客户可见区 A–H；K–O 为内部核算列（模板内自带公式）
DEFAULT_TEMPLATE_NAME = "proposal_template.xlsx"
# 店铺名（或关键词）→ 模板文件；按最长匹配优先
SHOP_TEMPLATE_MAP: tuple[tuple[str, str], ...] = (
    ("通江东晖电子商务有限公司", "tjdh_template.xlsx"),
    ("云上（北川）人工智能科技有限公司", "ysbc_template.xlsx"),
    ("广元阡陌农业发展有限公司", "gyqm_template.xlsx"),
    ("拾味山林（广元）农业有限公司", "swsl_template.xlsx"),
    ("宣汉谷满田园农业有限公司", "xhyp_template.xlsx"),
)

# 商品行公式（相对模板第 3 行写法；合计行另处理）
_PRODUCT_FORMULAS = {
    7: "=E{row}*F{row}",  # 优惠合计
    12: "=G{row}*0.08",  # 运费
    13: "=K{row}*F{row}",  # 总成本
    14: "=G{row}-M{row}",  # 总利润
    15: "=N{row}/G{row}",  # 毛利率
}

_CELL_REF_RE = re.compile(r"(\$?)([A-Z]+)(\$?)(\d+)")
_DISPIMG_RE = re.compile(r'DISPIMG\("([^"]+)"', re.I)
# 图片与单元格边框之间留出的像素，避免盖住分割线
_IMAGE_PAD_PX = 4


def _backend_data_dir() -> Path:
    return Path(__file__).resolve().parents[2] / "data"


def proposal_template_path(name: str | None = None) -> Path:
    configured = (os.getenv("PROPOSAL_TEMPLATE_PATH") or "").strip()
    if configured and not name:
        return Path(configured)
    return _backend_data_dir() / (name or DEFAULT_TEMPLATE_NAME)


def _normalize_shop(value: str) -> str:
    text = (value or "").strip().lower()
    for token in (" ", "\u3000", "（", "）", "(", ")"):
        text = text.replace(token, "")
    return text


def _single_supplier_name(spec: dict) -> str:
    """只有全部商品都来自同一家店铺时，才返回该店铺名。"""
    shops = {
        str(line.get("supplier_name") or "").strip()
        for line in (spec.get("lines") or [])
        if str(line.get("supplier_name") or "").strip()
    }
    if len(shops) == 1:
        return next(iter(shops))
    return ""


def resolve_template_name(spec: dict) -> str:
    """仅当商品全部来自同一家店铺时，才使用专属模板。"""
    single_shop = _single_supplier_name(spec)
    if not single_shop:
        return DEFAULT_TEMPLATE_NAME
    ranked = sorted(SHOP_TEMPLATE_MAP, key=lambda item: len(item[0]), reverse=True)
    normalized_hint = _normalize_shop(single_shop)
    if not normalized_hint:
        return DEFAULT_TEMPLATE_NAME
    for shop, filename in ranked:
        normalized_shop = _normalize_shop(shop)
        if (
            normalized_hint == normalized_shop
            or normalized_hint in normalized_shop
            or normalized_shop in normalized_hint
        ):
            return filename
    return DEFAULT_TEMPLATE_NAME


def _header_text(value) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def _template_is_usable(path: Path) -> tuple[bool, str]:
    if not path.is_file():
        return False, "文件不存在"
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            sheet = workbook.active
            headers = [_header_text(sheet.cell(2, column).value) for column in range(1, 9)]
            if "名称" not in headers[0]:
                return False, f"A2 不是产品名称列：{headers[0]!r}"
            joined = "".join(headers)
            soft_missing = [
                token for token in ("规格", "平台", "优惠", "数量") if token not in joined
            ]
            if soft_missing:
                return False, f"第2行表头缺少关键列：{soft_missing} actual={headers}"
            if sheet.max_row < 4:
                return False, "缺少样例数据行或合计区"
            return True, ""
        finally:
            workbook.close()
    except Exception as error:
        return False, str(error)


def _find_total_row(sheet) -> int | None:
    for row in range(3, (sheet.max_row or 3) + 1):
        value = sheet.cell(row, 1).value
        if isinstance(value, str) and "合计" in value.replace(" ", ""):
            return row
    return None


def _snapshot_row_styles(sheet, row: int, columns: int = 15) -> list:
    return [copy(sheet.cell(row, column)._style) for column in range(1, columns + 1)]


def _snapshot_footer(sheet, start_row: int) -> list[dict]:
    """合计行之后的说明行：值、样式、行高、涉及该行的合并。"""
    footers: list[dict] = []
    max_row = sheet.max_row or start_row
    merges = [
        str(merged)
        for merged in sheet.merged_cells.ranges
        if merged.min_row >= start_row
    ]
    for row in range(start_row, max_row + 1):
        values = [sheet.cell(row, column).value for column in range(1, 16)]
        if all(value is None for value in values):
            continue
        footers.append(
            {
                "values": values,
                "styles": _snapshot_row_styles(sheet, row),
                "height": sheet.row_dimensions[row].height,
            }
        )
    return [{"merges": merges, "rows": footers}]


def _clear_sheet_images(sheet) -> None:
    images = getattr(sheet, "_images", None)
    if images is not None:
        images.clear()


def _xml_local(tag: str) -> str:
    return tag.split("}")[-1] if "}" in tag else tag


def _load_cellimage_bytes(template: Path) -> dict[str, bytes]:
    """WPS cellimages.xml：DISPIMG 的 ID → 图片字节。openpyxl 保存时不会保留这份私有数据。"""
    mapping: dict[str, bytes] = {}
    try:
        with ZipFile(template) as archive:
            names = archive.namelist()
            if "xl/cellimages.xml" not in names:
                return mapping
            rels: dict[str, str] = {}
            if "xl/_rels/cellimages.xml.rels" in names:
                root = ET.fromstring(archive.read("xl/_rels/cellimages.xml.rels"))
                for node in root:
                    rel_id = node.attrib.get("Id")
                    target = (node.attrib.get("Target") or "").replace("\\", "/")
                    if rel_id and target and target.upper() != "NULL":
                        rels[rel_id] = target
            root = ET.fromstring(archive.read("xl/cellimages.xml"))
            for image_node in root:
                if _xml_local(image_node.tag) != "cellImage":
                    continue
                image_id = ""
                embed = ""
                for node in image_node.iter():
                    local = _xml_local(node.tag)
                    if local == "cNvPr":
                        image_id = node.attrib.get("name") or ""
                    elif local == "blip":
                        embed = (
                            node.attrib.get(
                                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                            )
                            or node.attrib.get("embed")
                            or ""
                        )
                target = rels.get(embed) or ""
                if not image_id or not target:
                    continue
                zip_name = target if target.startswith("xl/") else f"xl/{target.lstrip('/')}"
                if zip_name not in names:
                    continue
                mapping[image_id] = archive.read(zip_name)
    except Exception as error:
        logger.warning("读取模板 cellimages 失败 path={} error={}", template, error)
    return mapping


def _dispimg_id(value) -> str | None:
    if not isinstance(value, str) or "DISPIMG" not in value.upper():
        return None
    match = _DISPIMG_RE.search(value)
    return match.group(1) if match else None


def _image_bytes_from_drawing(image: ExcelImage) -> bytes | None:
    data = getattr(image, "_data", None)
    try:
        raw = data() if callable(data) else data
    except Exception:
        return None
    return raw if isinstance(raw, (bytes, bytearray)) else None


def _snapshot_header_drawings(sheet) -> list[dict]:
    """标题行（Excel 第 1 行）上的标准 drawing，避免清商品图时把 logo 一起删掉。"""
    kept: list[dict] = []
    for image in list(getattr(sheet, "_images", []) or []):
        origin = getattr(image.anchor, "_from", None)
        if origin is None:
            continue
        row_idx = getattr(origin, "row", None)
        col_idx = getattr(origin, "col", None)
        if row_idx is None or int(row_idx) > 0:
            continue
        raw = _image_bytes_from_drawing(image)
        if not raw:
            continue
        item = {
            "bytes": bytes(raw),
            "col": int(col_idx or 0) + 1,
            "row": int(row_idx) + 1,
            "colOff": int(getattr(origin, "colOff", 0) or 0),
            "rowOff": int(getattr(origin, "rowOff", 0) or 0),
            "width": int(image.width or 0),
            "height": int(image.height or 0),
        }
        dest = getattr(image.anchor, "to", None)
        if dest is not None and getattr(dest, "col", None) is not None:
            item["to"] = {
                "col": int(dest.col) + 1,
                "row": int(dest.row) + 1,
                "colOff": int(getattr(dest, "colOff", 0) or 0),
                "rowOff": int(getattr(dest, "rowOff", 0) or 0),
            }
        kept.append(item)
    return kept


def _column_width_chars(sheet, column: int) -> float:
    letter = get_column_letter(column)
    width = sheet.column_dimensions[letter].width
    if width:
        return float(width)
    default = getattr(sheet.sheet_format, "defaultColWidth", None)
    return float(default or 8.43)


def _row_height_points(sheet, row: int) -> float:
    height = sheet.row_dimensions[row].height
    if height:
        return float(height)
    default = getattr(sheet.sheet_format, "defaultRowHeight", None)
    return float(default or 15)


def _merge_box(sheet, column: int, row: int) -> tuple[int, int, int, int]:
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
            return merged.min_col, merged.min_row, merged.max_col, merged.max_row
    return column, row, column, row


def _box_size_px(sheet, min_col: int, min_row: int, max_col: int, max_row: int) -> tuple[int, int]:
    width = sum(_column_width_chars(sheet, col) * 7 + 5 for col in range(min_col, max_col + 1))
    height = sum(points_to_pixels(_row_height_points(sheet, row)) for row in range(min_row, max_row + 1))
    return max(8, int(width)), max(8, int(height))


def _add_image(
    sheet,
    image: ExcelImage,
    *,
    column: int,
    row: int,
    offset_x: int,
    offset_y: int,
) -> None:
    image.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=column - 1,
            row=row - 1,
            colOff=pixels_to_EMU(max(0, offset_x)),
            rowOff=pixels_to_EMU(max(0, offset_y)),
        ),
        ext=XDRPositiveSize2D(
            cx=pixels_to_EMU(max(1, int(image.width))),
            cy=pixels_to_EMU(max(1, int(image.height))),
        ),
    )
    sheet.add_image(image)


def _fit_image_in_box(
    sheet,
    image: ExcelImage,
    *,
    min_col: int,
    min_row: int,
    max_col: int,
    max_row: int,
    pad_px: int = _IMAGE_PAD_PX,
) -> None:
    """按单元格（或合并区）等比缩放并居中，四周留白以免盖住分割线。"""
    box_w, box_h = _box_size_px(sheet, min_col, min_row, max_col, max_row)
    inner_w = max(8, box_w - pad_px * 2)
    inner_h = max(8, box_h - pad_px * 2)
    orig_w = max(1, int(image.width or 1))
    orig_h = max(1, int(image.height or 1))
    scale = min(inner_w / orig_w, inner_h / orig_h)
    image.width = max(1, int(orig_w * scale))
    image.height = max(1, int(orig_h * scale))
    offset_x = pad_px + max(0, (inner_w - int(image.width)) // 2)
    offset_y = pad_px + max(0, (inner_h - int(image.height)) // 2)
    _add_image(sheet, image, column=min_col, row=min_row, offset_x=offset_x, offset_y=offset_y)


def _excel_image_from_bytes(raw: bytes, buffers: list[io.BytesIO]) -> ExcelImage:
    with PillowImage.open(io.BytesIO(raw)) as image:
        if image.mode in ("RGBA", "LA") or (
            image.mode == "P" and "transparency" in image.info
        ):
            rgba = image.convert("RGBA")
            background = PillowImage.new("RGBA", rgba.size, (255, 255, 255, 255))
            converted = PillowImage.alpha_composite(background, rgba).convert("RGB")
        else:
            converted = image.convert("RGB")
        output = io.BytesIO()
        converted.save(output, format="PNG")
        width, height = converted.size
    output.seek(0)
    buffers.append(output)
    excel_image = ExcelImage(output)
    excel_image.width = width
    excel_image.height = height
    return excel_image


def _restore_header_logos(
    sheet,
    *,
    drawings: list[dict],
    dispimg_bytes: dict[str, bytes],
    buffers: list[io.BytesIO],
    template_name: str = "",
) -> None:
    occupied: set[tuple[int, int]] = set()
    tpl = template_name.lower()

    for item in drawings:
        excel_image = _excel_image_from_bytes(item["bytes"], buffers)
        dest = item.get("to")
        if dest:
            # TwoCellAnchor 按模板原锚点还原，避免用原图像素把 logo 撑破
            excel_image.anchor = TwoCellAnchor(
                _from=AnchorMarker(
                    col=item["col"] - 1,
                    row=item["row"] - 1,
                    colOff=int(item["colOff"] or 0),
                    rowOff=int(item["rowOff"] or 0),
                ),
                to=AnchorMarker(
                    col=dest["col"] - 1,
                    row=dest["row"] - 1,
                    colOff=int(dest["colOff"] or 0),
                    rowOff=int(dest["rowOff"] or 0),
                ),
            )
            sheet.add_image(excel_image)
        else:
            if item["width"] > 0 and item["height"] > 0:
                excel_image.width = item["width"]
                excel_image.height = item["height"]
            offset_x = int(round((item["colOff"] or 0) / 9525))
            offset_y = int(round((item["rowOff"] or 0) / 9525))
            if "tjdh" in tpl:
                row_h = _row_height_points(sheet, item["row"])
                offset_y = max(offset_y, int(points_to_pixels(row_h) * 0.20))
            _add_image(
                sheet,
                excel_image,
                column=item["col"],
                row=item["row"],
                offset_x=offset_x,
                offset_y=offset_y,
            )
        occupied.add((item["col"], item["row"]))

    for column in range(1, 9):
        cell = sheet.cell(1, column)
        image_id = _dispimg_id(cell.value)
        if not image_id or image_id not in dispimg_bytes:
            continue
        if (column, 1) in occupied:
            cell.value = None
            continue
        excel_image = _excel_image_from_bytes(dispimg_bytes[image_id], buffers)
        if "ysbc" in tpl and column == 1:
            # 云上北川：logo 拉伸铺满 A1:H1，用 TwoCellAnchor 确保右侧齐平
            excel_image.anchor = TwoCellAnchor(
                _from=AnchorMarker(col=0, row=0, colOff=0, rowOff=0),
                to=AnchorMarker(col=8, row=1, colOff=0, rowOff=0),
            )
            sheet.add_image(excel_image)
        else:
            min_col, min_row, max_col, max_row = _merge_box(sheet, column, 1)
            _fit_image_in_box(
                sheet,
                excel_image,
                min_col=min_col,
                min_row=min_row,
                max_col=max_col,
                max_row=max_row,
                pad_px=0,
            )
        cell.value = None
        occupied.add((column, 1))


def _shift_product_formula(formula: str, dst_row: int) -> str:
    """商品行公式：样例在第 3 行，把其中对第 3 行的引用改到 dst_row。"""

    def repl(match: re.Match[str]) -> str:
        abs_col, col, abs_row, row_s = match.groups()
        row = int(row_s)
        if row == 3:
            return f"{abs_col}{col}{abs_row}{dst_row}"
        return match.group(0)

    return _CELL_REF_RE.sub(repl, formula)


def _shift_total_formula(
    formula: str, *, sample_total_row: int, total_row: int, last_data: int
) -> str:
    """合计行公式：自身行号跟合计行走；数据区结束行改到 last_data。"""
    sample_last = sample_total_row - 1

    def repl(match: re.Match[str]) -> str:
        abs_col, col, abs_row, row_s = match.groups()
        row = int(row_s)
        if row == sample_total_row:
            return f"{abs_col}{col}{abs_row}{total_row}"
        if row == sample_last:
            return f"{abs_col}{col}{abs_row}{last_data}"
        return match.group(0)

    return _CELL_REF_RE.sub(repl, formula)


def _product_cell_value(column: int, row: int, sample, fallback):
    if column == 2:
        return None  # 图片列改用 drawing，去掉 DISPIMG
    if isinstance(sample, str) and sample.startswith("="):
        if "DISPIMG" in sample.upper():
            return None
        return _shift_product_formula(sample, row)
    if column in _PRODUCT_FORMULAS:
        return _PRODUCT_FORMULAS[column].format(row=row)
    return fallback


def _total_cell_value(column: int, sample, *, sample_total_row: int, total_row: int, last_data: int):
    if not (isinstance(sample, str) and sample.startswith("=")):
        return sample
    letter = get_column_letter(column)
    compact = sample.replace(" ", "").upper()
    # 数据区纵向合计：无论原样是 D3:D5 还是 D3+D4+D5，都按实际行数重写
    range_sum = re.fullmatch(rf"=SUM\({letter}\d+:{letter}\d+\)", compact)
    plus_sum = re.fullmatch(rf"=SUM\({letter}\d+(\+{letter}\d+)+\)", compact)
    if range_sum or plus_sum:
        return f"=SUM({letter}3:{letter}{last_data})"
    return _shift_total_formula(
        sample,
        sample_total_row=sample_total_row,
        total_row=total_row,
        last_data=last_data,
    )


def _place_product_image(sheet, image: ExcelImage, column: int, row: int) -> None:
    min_col, min_row, max_col, max_row = _merge_box(sheet, column, row)
    _fit_image_in_box(
        sheet,
        image,
        min_col=min_col,
        min_row=min_row,
        max_col=max_col,
        max_row=max_row,
        pad_px=_IMAGE_PAD_PX,
    )


def _absolute_image_url(raw: str) -> str:
    value = (raw or "").strip()
    if value.startswith(("http://", "https://")):
        return value
    base = (os.getenv("PUBLIC_BASE_URL") or os.getenv("API_PUBLIC_BASE_URL") or "").strip()
    if base and value:
        return urljoin(base.rstrip("/") + "/", value.lstrip("/"))
    return value


async def _download_image(url: str) -> bytes | None:
    resolved = _absolute_image_url(url)
    if not resolved:
        return None
    if resolved.startswith("/media/"):
        backend_dir = Path(__file__).resolve().parents[2]
        local_path = backend_dir / resolved.lstrip("/")
        try:
            return local_path.read_bytes()
        except OSError:
            return None
    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(12.0, connect=5.0)) as client:
            response = await client.get(resolved, follow_redirects=True)
            response.raise_for_status()
            if len(response.content) > 8 * 1024 * 1024:
                return None
            return response.content
    except Exception:
        return None


def _placeholder_bytes() -> bytes:
    image = PillowImage.new("RGB", (240, 180), color=(242, 244, 247))
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _normalized_png(raw: bytes | None) -> bytes:
    try:
        with PillowImage.open(io.BytesIO(raw or b"")) as image:
            image.thumbnail((1200, 1200))
            if image.mode in ("RGBA", "LA") or (
                image.mode == "P" and "transparency" in image.info
            ):
                rgba = image.convert("RGBA")
                background = PillowImage.new("RGBA", rgba.size, (255, 255, 255, 255))
                converted = PillowImage.alpha_composite(background, rgba).convert("RGB")
            else:
                converted = image.convert("RGB")
            output = io.BytesIO()
            converted.save(output, format="PNG", optimize=True)
            return output.getvalue()
    except Exception:
        return _placeholder_bytes()


def _insert_soft_wrap(text, width: int = 10) -> str | None:
    """WPS 对 wrap_text 支持不稳定，备注/名称过长时插入硬换行。"""
    if text is None:
        return None
    raw = str(text).strip()
    if not raw or "\n" in raw or len(raw) <= width:
        return raw
    for sep in ("，", "、", "；", ";", " ", ",", "。"):
        idx = raw.find(sep)
        if 4 <= idx <= len(raw) - 3:
            return raw[: idx + 1] + "\n" + raw[idx + 1 :].lstrip()
    mid = max(width, (len(raw) + 1) // 2)
    return raw[:mid] + "\n" + raw[mid:]


def _force_wrap(cell) -> None:
    current = cell.alignment
    cell.alignment = Alignment(
        horizontal=current.horizontal or "center",
        vertical=current.vertical or "center",
        wrap_text=True,
        shrink_to_fit=current.shrink_to_fit,
        indent=current.indent,
        text_rotation=current.text_rotation,
    )


def _apply_styles(sheet, row: int, styles: list) -> None:
    for column, style in enumerate(styles, 1):
        sheet.cell(row, column)._style = copy(style)


async def _render_template_xlsx(spec: dict, target: Path, template: Path) -> None:
    workbook = load_workbook(template, data_only=False)
    sheet = workbook.active
    sample_total = _find_total_row(sheet)
    if sample_total is None:
        workbook.close()
        raise ValueError("模板缺少合计行")

    product_style = _snapshot_row_styles(sheet, 3)
    next_style = (
        _snapshot_row_styles(sheet, 4)
        if sample_total > 4
        else product_style
    )
    total_styles = _snapshot_row_styles(sheet, sample_total)
    sample_product_formulas = {
        column: sheet.cell(3, column).value for column in range(1, 16)
    }
    sample_total_values = {
        column: sheet.cell(sample_total, column).value for column in range(1, 16)
    }
    product_height = sheet.row_dimensions[3].height or 100
    total_height = sheet.row_dimensions[sample_total].height
    footer_pack = _snapshot_footer(sheet, sample_total + 1)[0]
    total_merges = [
        str(merged)
        for merged in sheet.merged_cells.ranges
        if merged.min_row == sample_total and merged.max_row == sample_total
    ]
    header_drawings = _snapshot_header_drawings(sheet)
    dispimg_bytes = _load_cellimage_bytes(template)

    # 保留标题与表头；清掉样例商品/合计/页脚。模板图是 WPS DISPIMG，改为 drawing。
    for merged in list(sheet.merged_cells.ranges):
        if merged.min_row >= 3:
            sheet.unmerge_cells(str(merged))
    _clear_sheet_images(sheet)
    if sheet.max_row >= 3:
        sheet.delete_rows(3, sheet.max_row - 2)

    lines = spec.get("lines") or []
    image_buffers: list[io.BytesIO] = []
    for offset, line in enumerate(lines):
        row_index = 3 + offset
        styles = product_style if offset % 2 == 0 else next_style
        # A 名称 / B 图 / C 规格 / D 平台价 / E 优惠价 / F 数量 / G 公式 / H 备注 / K 成本 / L-O 公式
        fallbacks = {
            1: _insert_soft_wrap(line.get("product_name"), 8),
            2: None,
            3: line.get("spec"),
            4: line.get("platform_price"),
            5: line.get("promo_unit_price"),
            6: line.get("qty"),
            8: _insert_soft_wrap(line.get("remark"), 10),
            11: line.get("cost_price"),
        }
        for column in range(1, 16):
            value = _product_cell_value(
                column,
                row_index,
                sample_product_formulas.get(column),
                fallbacks.get(column),
            )
            cell = sheet.cell(row_index, column, value)
            cell._style = copy(styles[column - 1])
            if column in (1, 3, 8):
                _force_wrap(cell)
            if column in (4, 5, 7, 11, 12, 13, 14):
                cell.number_format = "¥0.00"
            if column == 15:
                cell.number_format = "0.00%"
        sheet.row_dimensions[row_index].height = product_height

        image_data = _normalized_png(
            await _download_image(str(line.get("image_url") or ""))
        )
        with PillowImage.open(io.BytesIO(image_data)) as preview:
            orig_w, orig_h = preview.size
        buffer = io.BytesIO(image_data)
        image_buffers.append(buffer)
        excel_image = ExcelImage(buffer)
        excel_image.width = orig_w
        excel_image.height = orig_h
        _place_product_image(sheet, excel_image, 2, row_index)

    last_data = 3 + len(lines) - 1 if lines else 3
    total_row = last_data + 1 if lines else 3
    _apply_styles(sheet, total_row, total_styles)
    if total_height:
        sheet.row_dimensions[total_row].height = total_height

    for column in range(1, 16):
        sample_value = sample_total_values.get(column)
        if column == 1:
            value = "合计"
        elif column == 10 and isinstance(sample_value, str) and "合计" in sample_value:
            value = "合计"
        else:
            value = _total_cell_value(
                column,
                sample_value,
                sample_total_row=sample_total,
                total_row=total_row,
                last_data=last_data,
            )
        cell = sheet.cell(total_row, column, value)
        cell._style = copy(total_styles[column - 1])
        if column in (4, 5, 7, 11, 12, 13, 14):
            cell.number_format = "¥0.00"
        if column == 15:
            cell.number_format = "0.00%"

    for merge in total_merges:
        # 形如 B5:C5 → 挪到合计行
        match = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)$", merge)
        if not match:
            continue
        c1, _, c2, _ = match.groups()
        sheet.merge_cells(f"{c1}{total_row}:{c2}{total_row}")

    footer_start = total_row + 1
    for offset, footer in enumerate(footer_pack["rows"]):
        row_index = footer_start + offset
        for column, value in enumerate(footer["values"], 1):
            cell = sheet.cell(row_index, column, value)
            cell._style = copy(footer["styles"][column - 1])
        if footer.get("height"):
            sheet.row_dimensions[row_index].height = footer["height"]
    for merge in footer_pack["merges"]:
        match = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)$", merge)
        if not match:
            continue
        c1, r1, c2, r2 = match.groups()
        delta = footer_start - (sample_total + 1)
        sheet.merge_cells(f"{c1}{int(r1) + delta}:{c2}{int(r2) + delta}")

    _restore_header_logos(
        sheet,
        drawings=header_drawings,
        dispimg_bytes=dispimg_bytes,
        buffers=image_buffers,
        template_name=template.name,
    )

    sheet.sheet_view.showGridLines = False
    end_row = footer_start + max(len(footer_pack["rows"]) - 1, 0)
    sheet.print_area = f"A1:O{end_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    workbook.close()


async def _render_generated_xlsx(spec: dict, target: Path) -> None:
    """无可用模板时的代码生成回退：列布局与新模板一致，并写入计算函数。"""
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "方案"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:H1")
    title_cell = sheet["A1"]
    title_cell.value = spec.get("title") or "脱贫地区农副产品网络销售平台 产品供应表（包邮）"
    title_cell.font = Font(name="微软雅黑", size=18, bold=True)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = PatternFill("solid", fgColor="DDEBF7")
    sheet.row_dimensions[1].height = 34

    headers = [
        "产品名称",
        "产品图片",
        "规格",
        "平台价格",
        "优惠价",
        "数量",
        "优惠合计",
        "备注",
        "",
        "",
        "成本",
        "运费",
        "总成本",
        "总利润",
        "毛利率",
    ]
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(2, column, header or None)
        cell.font = Font(name="微软雅黑", bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    image_buffers: list[io.BytesIO] = []
    lines = spec.get("lines") or []
    for row_index, line in enumerate(lines, 3):
        values = {
            1: line.get("product_name"),
            3: line.get("spec"),
            4: line.get("platform_price"),
            5: line.get("promo_unit_price"),
            6: line.get("qty"),
            7: f"=E{row_index}*F{row_index}",
            8: line.get("remark"),
            11: line.get("cost_price"),
            12: f"=G{row_index}*0.08",
            13: f"=K{row_index}*F{row_index}",
            14: f"=G{row_index}-M{row_index}",
            15: f"=N{row_index}/G{row_index}",
        }
        for column in range(1, 16):
            cell = sheet.cell(row_index, column, values.get(column))
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if column in (4, 5, 7, 11, 12, 13, 14):
                cell.number_format = "¥0.00"
            if column == 15:
                cell.number_format = "0.00%"
        sheet.row_dimensions[row_index].height = 92
        image_data = _normalized_png(await _download_image(str(line.get("image_url") or "")))
        with PillowImage.open(io.BytesIO(image_data)) as preview:
            orig_w, orig_h = preview.size
        buffer = io.BytesIO(image_data)
        image_buffers.append(buffer)
        excel_image = ExcelImage(buffer)
        excel_image.width = orig_w
        excel_image.height = orig_h
        _place_product_image(sheet, excel_image, 2, row_index)

    last_data = 2 + len(lines)
    total_row = last_data + 1
    sheet.cell(total_row, 1, "合计")
    for column in (4, 5, 6, 7, 11, 12, 13, 14):
        letter = get_column_letter(column)
        sheet.cell(total_row, column, f"=SUM({letter}3:{letter}{last_data})")
        if column != 6:
            sheet.cell(total_row, column).number_format = "¥0.00"
    sheet.cell(total_row, 10, "合计")
    sheet.cell(total_row, 15, f"=(N{total_row}-L{total_row})/G{total_row}")
    sheet.cell(total_row, 15).number_format = "0.00%"
    for column in range(1, 16):
        sheet.cell(total_row, column).border = border
        sheet.cell(total_row, column).alignment = Alignment(
            horizontal="center", vertical="center"
        )
    sheet.merge_cells(start_row=total_row, start_column=2, end_row=total_row, end_column=3)

    widths = [18, 18, 12, 12, 12, 10, 12, 20, 3, 8, 10, 10, 10, 10, 10]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.print_area = f"A1:O{total_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    workbook.save(target)
    workbook.close()


async def render_xlsx(spec: dict, target: Path) -> None:
    template_name = resolve_template_name(spec)
    template = proposal_template_path(template_name)
    usable, reason = _template_is_usable(template)
    if not usable and template_name != DEFAULT_TEMPLATE_NAME:
        logger.warning(
            "店铺模板不可用，回退通用模板 path={} reason={}", template, reason
        )
        template = proposal_template_path(DEFAULT_TEMPLATE_NAME)
        usable, reason = _template_is_usable(template)
    if usable:
        try:
            await _render_template_xlsx(spec, target, template)
            logger.info(
                "方案 Excel 已使用模板 path={} shop={}",
                template,
                (spec.get("meta") or {}).get("primary_shop") or "",
            )
            return
        except Exception as error:
            logger.exception(
                "方案模板渲染失败，回退代码生成 path={} error={}", template, error
            )
    else:
        logger.warning("方案模板不可用，回退代码生成 path={} reason={}", template, reason)
    await _render_generated_xlsx(spec, target)
