"""从 XLSX 批量绑定主系统商品规格 ID（mibuddy_sp_id）。

表头兼容：id + 规格 ID 列（含历史「成本价」列，该列现为规格 ID）。
按本地主键 id 更新，空/#N/A 行跳过，不会清空已有规格 ID。
导入成功后立即按本次写入的 ID 刷新成本价。
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select

from core.product_cost_sync import sync_cost_prices_from_mibuddy
from core.product_sp_id import parse_mibuddy_sp_id
from models import Product

_ID_HEADERS = frozenset({"id", "商品id", "本地id", "主键"})
_SP_ID_HEADERS = frozenset(
    {
        "成本价",
        "cost_price",
        "成本",
        "成本价(元)",
        "规格id",
        "规格id(sp_id)",
        "spid",
        "sp_id",
        "mibuddy_sp_id",
        "ids",
        "商品规格id",
        "主系统规格id",
    }
)


def _header_key(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "").replace("（", "(").replace("）", ")")


def _find_columns(headers: list[Any]) -> tuple[int | None, int | None]:
    id_headers = {_header_key(h) for h in _ID_HEADERS}
    sp_headers = {_header_key(h) for h in _SP_ID_HEADERS}
    id_col = sp_col = None
    for index, raw in enumerate(headers):
        key = _header_key(raw)
        if id_col is None and key in id_headers:
            id_col = index
        if sp_col is None and key in sp_headers:
            sp_col = index
    return id_col, sp_col


def _empty_stats(*, errors: list[str] | None = None) -> dict[str, Any]:
    return {
        "updated": 0,
        "skipped": 0,
        "not_found": 0,
        "invalid": 0,
        "cost_updated": 0,
        "cost_not_found": 0,
        "cost_failed_batches": 0,
        "missing_sp_ids": [],
        "errors": errors or [],
    }


async def import_sp_ids_from_xlsx(db, content: bytes) -> dict[str, Any]:
    """解析并写入规格 ID，随后刷新对应成本价。"""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration:
            return _empty_stats(errors=["文件为空或没有表头"])
        id_col, sp_col = _find_columns(headers)
        if id_col is None or sp_col is None:
            return _empty_stats(
                errors=[
                    "未识别表头：需要同时包含「id」与规格 ID 列"
                    "（兼容「成本价」「sp_id」「mibuddy_sp_id」「ids」等）。"
                    f" 当前表头：{[str(h or '') for h in headers]}"
                ]
            )

        updates: list[tuple[int, str, int]] = []
        skipped = invalid = 0
        errors: list[str] = []
        seen_product_ids: set[int] = set()
        for row_index, row in enumerate(rows, start=2):
            if not row or all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            raw_id = row[id_col] if id_col < len(row) else None
            raw_sp = row[sp_col] if sp_col < len(row) else None
            try:
                product_id = int(raw_id)
            except (TypeError, ValueError):
                invalid += 1
                if len(errors) < 20:
                    errors.append(f"第 {row_index} 行：id 无效（{raw_id!r}）")
                continue
            sp_id, parse_error = parse_mibuddy_sp_id(raw_sp)
            if parse_error:
                invalid += 1
                if len(errors) < 20:
                    errors.append(f"第 {row_index} 行：{parse_error}")
                continue
            if sp_id is None:
                skipped += 1
                continue
            if product_id in seen_product_ids:
                skipped += 1
                continue
            seen_product_ids.add(product_id)
            updates.append((product_id, sp_id, row_index))
    finally:
        workbook.close()

    if not updates:
        return {
            **_empty_stats(errors=errors or ["没有可导入的规格 ID 行"]),
            "skipped": skipped,
            "invalid": invalid,
        }

    product_ids = [item[0] for item in updates]
    result = await db.execute(select(Product).where(Product.id.in_(product_ids)))
    by_id = {row.id: row for row in result.scalars().all()}

    # 同一规格 ID 允许绑到多条本地商品（商品库存在重复商品）
    updated = not_found = 0
    written_sp_ids: list[str] = []
    for product_id, sp_id, row_index in updates:
        product = by_id.get(product_id)
        if product is None:
            not_found += 1
            if len(errors) < 20:
                errors.append(f"第 {row_index} 行：id={product_id} 在商品库中不存在")
            continue
        product.mibuddy_sp_id = sp_id
        written_sp_ids.append(sp_id)
        updated += 1

    await db.commit()

    cost_updated = cost_not_found = cost_failed = 0
    missing_sp_ids: list[dict[str, Any]] = []
    if written_sp_ids:
        cost_stats = await sync_cost_prices_from_mibuddy(
            db,
            sp_ids=written_sp_ids,
            persist_status=False,
        )
        cost_updated = int(cost_stats.get("updated") or 0)
        cost_not_found = int(cost_stats.get("not_found") or 0)
        cost_failed = int(cost_stats.get("failed_batches") or 0)
        missing_sp_ids = list(cost_stats.get("missing_sp_ids") or [])
        for item in (cost_stats.get("errors") or [])[:10]:
            if len(errors) < 20:
                errors.append(f"刷新成本：{item}")

    return {
        "updated": updated,
        "skipped": skipped,
        "not_found": not_found,
        "invalid": invalid,
        "cost_updated": cost_updated,
        "cost_not_found": cost_not_found,
        "cost_failed_batches": cost_failed,
        "missing_sp_ids": missing_sp_ids,
        "errors": errors,
    }
