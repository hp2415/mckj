"""从 XLSX 按订单单号回填 raw_orders.wechat_idx。

表格需同时包含 DDDH、WECHAT 两列（表头大小写不敏感，兼容「订单号」「微信号」等别名）。
- DDDH：匹配 raw_orders.dddh
- WECHAT：写入 raw_orders.wechat_idx（销售号 alias_name）

默认只更新库中已存在、且 wechat_idx 为空的订单；不会覆盖已有归属。
数据库连接写在本文件顶部（DB_HOST 等），不读 .env。

用法（在 backend 目录）：
  python -m sync.import_raw_order_wechat_idx D:/data/orders.xlsx --dry-run
  python -m sync.import_raw_order_wechat_idx D:/data/orders.xlsx
  python -m sync.import_raw_order_wechat_idx D:/data/orders.xlsx --overwrite
  python -m sync.import_raw_order_wechat_idx D:/data/orders.xlsx --insert-missing
"""
from __future__ import annotations

import argparse
import asyncio
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote_plus

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects.mysql import insert as mysql_insert
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.orm import sessionmaker

from models import RawOrder

# 本脚本自建连接，不读 .env / database.py
DB_HOST = "192.168.0.100"
DB_PORT = 3306
DB_USER = "root"
DB_PASSWORD = "MySQLPassword$"
DB_NAME = "ai_assistant_db"

DATABASE_URL = (
    f"mysql+aiomysql://{quote_plus(DB_USER)}:{quote_plus(DB_PASSWORD)}"
    f"@{DB_HOST}:{DB_PORT}/{DB_NAME}"
)


def _header_key(value: Any) -> str:
    return (
        str(value or "")
        .strip()
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("（", "(")
        .replace("）", ")")
    )


_DDDH_HEADERS = frozenset(
    _header_key(h)
    for h in ("dddh", "订单单号", "订单号", "单号", "order_no", "orderno")
)
_WECHAT_HEADERS = frozenset(
    _header_key(h)
    for h in (
        "wechat",
        "微信",
        "微信号",
        "微信id",
        "wechat_idx",
        "alias",
        "alias_name",
        "别名",
        "销售微信",
        "销售号",
    )
)

BATCH_SIZE = 500
REPORT_CAP = 30


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        if value == int(value):
            return str(int(value))
        return str(value).strip()
    text = str(value).strip()
    if text.endswith(".0") and text.replace(".", "", 1).isdigit():
        return text[:-2]
    return text


def _find_columns(headers: list[Any]) -> tuple[int | None, int | None]:
    dddh_col = wechat_col = None
    for index, raw in enumerate(headers):
        key = _header_key(raw)
        if dddh_col is None and key in _DDDH_HEADERS:
            dddh_col = index
        if wechat_col is None and key in _WECHAT_HEADERS:
            wechat_col = index
    return dddh_col, wechat_col


def load_xlsx_pairs(path: Path, sheet_name: str | None = None) -> tuple[list[tuple[str, str]], dict[str, int]]:
    """返回 (dddh, wechat) 列表（后行覆盖先行）以及解析统计。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration:
            raise ValueError("文件为空或没有表头") from None

        dddh_col, wechat_col = _find_columns(headers)
        if dddh_col is None or wechat_col is None:
            raise ValueError(
                "未识别表头：需要同时包含 DDDH 与 WECHAT 列"
                f"（当前表头：{[str(h or '') for h in headers]}）"
            )

        mapping: dict[str, str] = {}
        stats = {
            "file_rows": 0,
            "empty_dddh": 0,
            "empty_wechat": 0,
            "dup_dddh": 0,
            "float_dddh": 0,
        }
        for row in rows:
            if not row or all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            stats["file_rows"] += 1
            raw_dddh = row[dddh_col] if dddh_col < len(row) else None
            if isinstance(raw_dddh, float) and raw_dddh == raw_dddh and abs(raw_dddh) >= 1e15:
                stats["float_dddh"] += 1
            dddh = _cell_str(raw_dddh)
            wechat = _cell_str(row[wechat_col] if wechat_col < len(row) else None)
            if not dddh:
                stats["empty_dddh"] += 1
                continue
            if not wechat:
                stats["empty_wechat"] += 1
                continue
            if dddh in mapping:
                stats["dup_dddh"] += 1
            mapping[dddh] = wechat
        return list(mapping.items()), stats
    finally:
        workbook.close()


def _make_session_factory():
    engine = create_async_engine(
        DATABASE_URL,
        echo=False,
        pool_size=5,
        max_overflow=5,
        pool_pre_ping=True,
        pool_recycle=280,
        connect_args={"connect_timeout": 10},
    )
    factory = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    return engine, factory


async def apply_pairs(
    pairs: list[tuple[str, str]],
    session_factory,
    *,
    overwrite: bool,
    insert_missing: bool,
    dry_run: bool,
) -> dict[str, Any]:
    updated = skipped_existing = not_found = inserted = 0
    not_found_dddh: list[str] = []
    preview: list[dict[str, str]] = []

    async with session_factory() as db:
        for offset in range(0, len(pairs), BATCH_SIZE):
            chunk = pairs[offset : offset + BATCH_SIZE]
            dddh_list = [d for d, _ in chunk]
            wanted = dict(chunk)

            result = await db.execute(select(RawOrder).where(RawOrder.dddh.in_(dddh_list)))
            orders = list(result.scalars().all())
            found_dddh: set[str] = set()

            for order in orders:
                dddh = str(order.dddh or "").strip()
                wechat = wanted.get(dddh)
                if not dddh or wechat is None:
                    continue
                found_dddh.add(dddh)
                current = str(order.wechat_idx or "").strip()
                if current and not overwrite:
                    skipped_existing += 1
                    continue
                if current == wechat:
                    skipped_existing += 1
                    continue
                if len(preview) < REPORT_CAP:
                    preview.append(
                        {
                            "dddh": dddh,
                            "order_id": str(order.order_id or ""),
                            "from": current,
                            "to": wechat,
                            "action": "update",
                        }
                    )
                if not dry_run:
                    order.wechat_idx = wechat
                updated += 1

            missing = [d for d in dddh_list if d not in found_dddh]
            if insert_missing:
                for dddh in missing:
                    wechat = wanted[dddh]
                    if len(preview) < REPORT_CAP:
                        preview.append(
                            {
                                "dddh": dddh,
                                "order_id": dddh,
                                "from": "",
                                "to": wechat,
                                "action": "insert",
                            }
                        )
                    if not dry_run:
                        now = datetime.now()
                        stmt = mysql_insert(RawOrder).values(
                            order_id=dddh,
                            dddh=dddh,
                            wechat_idx=wechat,
                            imported_at=now,
                            raw_json=json.dumps(
                                {"dddh": dddh, "wechat_idx": wechat, "source": "xlsx_wechat_idx"},
                                ensure_ascii=False,
                            ),
                        )
                        update_fields: dict[str, Any] = {
                            "dddh": stmt.inserted.dddh,
                            "raw_json": stmt.inserted.raw_json,
                        }
                        if overwrite:
                            update_fields["wechat_idx"] = stmt.inserted.wechat_idx
                        stmt = stmt.on_duplicate_key_update(**update_fields)
                        await db.execute(stmt)
                    inserted += 1
            else:
                not_found += len(missing)
                remain = REPORT_CAP - len(not_found_dddh)
                if remain > 0:
                    not_found_dddh.extend(missing[:remain])

        if not dry_run:
            await db.commit()

    return {
        "updated": updated,
        "inserted": inserted,
        "skipped_existing": skipped_existing,
        "not_found": not_found,
        "not_found_sample": not_found_dddh,
        "preview": preview,
        "dry_run": dry_run,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="按 DDDH 匹配 raw_orders，将 WECHAT 写入 wechat_idx"
    )
    parser.add_argument("xlsx", help="xlsx 文件路径")
    parser.add_argument("--sheet", default=None, help="工作表名称，默认第一个")
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="覆盖已有 wechat_idx（默认只填空值）",
    )
    parser.add_argument(
        "--insert-missing",
        action="store_true",
        help="库中没有该 dddh 时插入占位订单（order_id=dddh）",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="只解析并预览，不写库",
    )
    return parser


async def _run(args: argparse.Namespace) -> int:
    path = Path(args.xlsx).expanduser().resolve()
    if not path.is_file():
        print(f"ERROR: 文件不存在：{path}", file=sys.stderr)
        return 1

    pairs, parse_stats = load_xlsx_pairs(path, args.sheet)
    print(f"数据库：{DB_USER}@{DB_HOST}:{DB_PORT}/{DB_NAME}")
    print(f"文件：{path}")
    print(
        "解析："
        f"有效行={len(pairs)} file_rows={parse_stats['file_rows']} "
        f"空DDDH={parse_stats['empty_dddh']} 空WECHAT={parse_stats['empty_wechat']} "
        f"重复DDDH(后行覆盖)={parse_stats['dup_dddh']}"
    )
    if parse_stats.get("float_dddh"):
        print(
            f"警告：有 {parse_stats['float_dddh']} 行 DDDH 是 Excel 数值。"
            "超过 15 位的单号可能已失真，请把该列设为文本后再导入。",
            file=sys.stderr,
        )
    if not pairs:
        print("没有可导入的行。")
        return 1

    engine, session_factory = _make_session_factory()
    try:
        stats = await apply_pairs(
            pairs,
            session_factory,
            overwrite=args.overwrite,
            insert_missing=args.insert_missing,
            dry_run=args.dry_run,
        )
    finally:
        await engine.dispose()
    action = "预览" if stats["dry_run"] else "已写入"
    print(
        f"{action}：updated={stats['updated']} inserted={stats['inserted']} "
        f"skipped_existing={stats['skipped_existing']} not_found={stats['not_found']}"
    )
    if stats["preview"]:
        print("样例：")
        for item in stats["preview"]:
            print(
                f"  [{item['action']}] dddh={item['dddh']} "
                f"order_id={item['order_id']} {item['from']!r} -> {item['to']!r}"
            )
    if stats["not_found_sample"]:
        print("未匹配到的 dddh 样例：")
        for dddh in stats["not_found_sample"]:
            print(f"  {dddh}")
        if stats["not_found"] > len(stats["not_found_sample"]):
            print(f"  ... 另有 {stats['not_found'] - len(stats['not_found_sample'])} 条")
    return 0


def main() -> None:
    args = build_parser().parse_args()
    try:
        code = asyncio.run(_run(args))
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
    sys.exit(code)


if __name__ == "__main__":
    main()
